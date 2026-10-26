
import csv
import io
import re
import unicodedata

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models import BusinessRegister, Taxpayer, User, Conversation, TaxAssessment
from app.routers.admin_dashboard import require_admin
from app.services.business_reconciliation_service import build_reconciliation
from fastapi.responses import StreamingResponse




router = APIRouter(
    prefix="/api/admin/business-register",
    tags=["Business Register"],
)


class BusinessMatchReviewRequest(BaseModel):
    action: str
    taxpayer_id: int | None = None
    review_notes: str | None = None


# ------------------------------------------------------------
# Serialization
# ------------------------------------------------------------

def serialize_business(business: BusinessRegister) -> dict:
    return {
        "id": business.id,
        "business_name": business.business_name,
        "normalized_name": business.normalized_name,
        "registration_number": business.registration_number,
        "tax_identifier": business.tax_identifier,
        "phone_number": business.phone_number,
        "email": business.email,
        "address": business.address,
        "city": business.city,
        "state": business.state,
        "source_name": business.source_name,
        "source_url": business.source_url,
        "discovered_at": business.discovered_at,
        "imported_at": business.imported_at,
        "tax_year": business.tax_year,
        "matched_taxpayer_id": business.matched_taxpayer_id,
        "match_confidence": (
            float(business.match_confidence)
            if business.match_confidence is not None
            else None
        ),
        "match_status": business.match_status,
        "verification_status": business.verification_status,
        "review_notes": business.review_notes,
        "is_active": business.is_active,
        "created_at": business.created_at,
        "updated_at": business.updated_at,
    }


# ------------------------------------------------------------
# Business register listing
# ------------------------------------------------------------

@router.get("")
def list_businesses(
    search: str | None = Query(default=None, max_length=150),
    tax_year: int | None = Query(default=None, ge=2000, le=2100),
    match_status: str | None = Query(default=None, max_length=40),
    verification_status: str | None = Query(default=None, max_length=40),
    active_only: bool = True,
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    query = db.query(BusinessRegister)

    if active_only:
        query = query.filter(BusinessRegister.is_active.is_(True))

    if search and search.strip():
        term = f"%{search.strip()}%"
        query = query.filter(
            or_(
                BusinessRegister.business_name.ilike(term),
                BusinessRegister.normalized_name.ilike(term),
                BusinessRegister.registration_number.ilike(term),
                BusinessRegister.tax_identifier.ilike(term),
                BusinessRegister.phone_number.ilike(term),
                BusinessRegister.email.ilike(term),
                BusinessRegister.city.ilike(term),
            )
        )

    if tax_year is not None:
        query = query.filter(BusinessRegister.tax_year == tax_year)

    if match_status:
        query = query.filter(
            BusinessRegister.match_status == match_status.strip().upper()
        )

    if verification_status:
        query = query.filter(
            BusinessRegister.verification_status
            == verification_status.strip().upper()
        )

    total = query.count()

    businesses = (
        query.order_by(BusinessRegister.business_name.asc())
        .offset(offset)
        .limit(limit)
        .all()
    )

    return {
        "total": total,
        "limit": limit,
        "offset": offset,
        "items": [serialize_business(item) for item in businesses],
    }


# ------------------------------------------------------------
# Business reconciliation
# ------------------------------------------------------------

@router.get("/reconciliation")
def reconcile_business_register(
    tax_year: int | None = Query(default=None, ge=2000, le=2100),
    limit: int = Query(default=50, ge=1, le=200),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    return build_reconciliation(
        db=db,
        tax_year=tax_year,
        limit=limit,
        offset=offset,
    )


FILED_STATUS = "TP_FILED"
NOT_FILED_STATUS = "TP_NOT_FILED"
ALLOWED_FILING_STATUSES = {FILED_STATUS, NOT_FILED_STATUS}

# These statuses mean the assessment progressed beyond calculation alone.
# presented_at is also checked because it records when the liability was shown.
PRESENTED_ASSESSMENT_STATUSES = {
    "PRESENTED_TO_TAXPAYER",
    "ACCEPTED",
    "PAYMENT_PENDING",
    "PAID",
    "OBJECTED",
    "CLOSED",
}


def _build_business_status_rows(
    db: Session,
    tax_year: int | None = None,
    search: str | None = None,
    filing_status: str | None = None,
) -> list[dict]:
    """
    Calculate the Business Register status without changing stored records.

    TP_FILED requires:
      1. A reliable, unambiguous taxpayer match.
      2. A completed conversation.
      3. A matching assessment for the relevant tax year.
      4. Evidence that the assessment was presented to the taxpayer.

    Payment is not required.
    """
    if filing_status:
        filing_status = filing_status.strip().upper()
        if filing_status not in ALLOWED_FILING_STATUSES:
            raise HTTPException(
                status_code=400,
                detail="filing_status must be TP_FILED or TP_NOT_FILED.",
            )

    # Reconcile the complete active register before applying status filters.
    reconciliation = build_reconciliation(
        db=db,
        tax_year=tax_year,
        limit=100000,
        offset=0,
    )
    reconciliation_items = reconciliation.get("items", [])

    business_ids = [
        item.get("business", {}).get("id")
        for item in reconciliation_items
        if item.get("business", {}).get("id") is not None
    ]

    if not business_ids:
        return []

    businesses = (
        db.query(BusinessRegister)
        .filter(BusinessRegister.id.in_(business_ids))
        .all()
    )
    business_by_id = {business.id: business for business in businesses}

    active_taxpayer_ids = {
        row.id
        for row in db.query(Taxpayer.id)
        .filter(Taxpayer.is_active.is_(True))
        .all()
    }

    # Retrieve completed conversations and their assessment records in one
    # batch. An assessment being CALCULATED alone is not enough.
    assessment_rows = (
        db.query(TaxAssessment, Conversation.status)
        .join(
            Conversation,
            Conversation.id == TaxAssessment.conversation_id,
        )
        .filter(
            TaxAssessment.taxpayer_id.in_(active_taxpayer_ids or {-1}),
            Conversation.status == "COMPLETED",
        )
        .all()
    )

    presented_returns: set[tuple[int, int]] = set()

    for assessment, conversation_status in assessment_rows:
        assessment_status = (
            assessment.calculation_status or ""
        ).strip().upper()

        liability_presented = (
            assessment.presented_at is not None
            or assessment_status in PRESENTED_ASSESSMENT_STATUSES
        )

        if (
            conversation_status == "COMPLETED"
            and liability_presented
        ):
            presented_returns.add(
                (assessment.taxpayer_id, assessment.tax_year)
            )

    output = []

    for item in reconciliation_items:
        business_data = item.get("business", {})
        business_id = business_data.get("id")
        business = business_by_id.get(business_id)

        if business is None:
            continue

        target_year = (
            business.tax_year
            if business.tax_year is not None
            else tax_year
        )

        candidates = item.get("candidate_matches", [])

        # Only a unique highest-scoring candidate with a score of at least
        # 80 is automatically considered a reliable match.
        scored_candidates = []
        for candidate in candidates:
            taxpayer_id = candidate.get("taxpayer_id")
            score = candidate.get(
                "confidence_score",
                candidate.get("confidence", 0),
            )

            try:
                score = int(score or 0)
            except (TypeError, ValueError):
                score = 0

            if taxpayer_id in active_taxpayer_ids:
                scored_candidates.append((taxpayer_id, score))

        best_score = max(
            (score for _, score in scored_candidates),
            default=0,
        )
        best_candidates = [
            taxpayer_id
            for taxpayer_id, score in scored_candidates
            if score == best_score
        ]

        reliable_taxpayer_id = None

        # A previously confirmed match is accepted only while its taxpayer
        # still exists and is active.
        if (
            business.match_status
            and business.match_status.strip().upper() == "CONFIRMED"
            and business.matched_taxpayer_id in active_taxpayer_ids
        ):
            reliable_taxpayer_id = business.matched_taxpayer_id

        elif best_score >= 80 and len(set(best_candidates)) == 1:
            reliable_taxpayer_id = best_candidates[0]

        return_completed = False

        if reliable_taxpayer_id is not None:
            return_completed = any(
                taxpayer_id == reliable_taxpayer_id
                and (
                    target_year is None
                    or assessment_year == target_year
                )
                for taxpayer_id, assessment_year in presented_returns
            )

        status = (
            FILED_STATUS
            if reliable_taxpayer_id is not None and return_completed
            else NOT_FILED_STATUS
        )

        row = {
            **business_data,
            "tax_year": business.tax_year,
            "matched_taxpayer_id": reliable_taxpayer_id,
            "filing_status": status,
            "filing_status_label": (
                "TP Filed" if status == FILED_STATUS else "TP Not Filed"
            ),
        }

        if search and search.strip():
            term = search.strip().casefold()
            searchable_values = (
                business.business_name,
                business.registration_number,
                business.tax_identifier,
                business.phone_number,
                business.email,
                business.city,
            )
            if not any(
                term in str(value or "").casefold()
                for value in searchable_values
            ):
                continue

        if filing_status and status != filing_status:
            continue

        output.append(row)

    return output


@router.get("/status-report")
def business_status_report(
    search: str | None = Query(default=None, max_length=150),
    tax_year: int | None = Query(default=None, ge=2000, le=2100),
    filing_status: str | None = Query(default=None),
    limit: int = Query(default=50, ge=1, le=500),
    offset: int = Query(default=0, ge=0),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    rows = _build_business_status_rows(
        db=db,
        tax_year=tax_year,
        search=search,
        filing_status=filing_status,
    )

    return {
        "total": len(rows),
        "limit": limit,
        "offset": offset,
        "items": rows[offset:offset + limit],
    }


@router.get("/export.csv")
def export_business_status_report(
    search: str | None = Query(default=None, max_length=150),
    tax_year: int | None = Query(default=None, ge=2000, le=2100),
    filing_status: str | None = Query(default=None),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    rows = _build_business_status_rows(
        db=db,
        tax_year=tax_year,
        search=search,
        filing_status=filing_status,
    )

    buffer = io.StringIO()
    columns = [
        "business_name",
        "registration_number",
        "tax_identifier",
        "phone_number",
        "email",
        "address",
        "city",
        "state",
        "tax_year",
        "filing_status_label",
    ]

    writer = csv.DictWriter(buffer, fieldnames=columns, extrasaction="ignore")
    writer.writeheader()
    writer.writerows(rows)

    filename = "eirs_business_register.csv"
    return StreamingResponse(
        iter([buffer.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"'
        },
    )

# ------------------------------------------------------------
# Business record detail
# ------------------------------------------------------------

@router.get("/{business_id}")
def get_business(
    business_id: int,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    business = (
        db.query(BusinessRegister)
        .filter(BusinessRegister.id == business_id)
        .first()
    )

    if business is None:
        raise HTTPException(
            status_code=404,
            detail="Business register record not found.",
        )

    result = serialize_business(business)

    result["matched_taxpayer"] = None

    if business.matched_taxpayer_id is not None:
        taxpayer = (
            db.query(Taxpayer)
            .filter(Taxpayer.id == business.matched_taxpayer_id)
            .first()
        )

        if taxpayer:
            result["matched_taxpayer"] = {
                "id": taxpayer.id,
                "taxpayer_reference": taxpayer.taxpayer_reference,
                "full_name": taxpayer.full_name,
            }

    return result


# ------------------------------------------------------------
# Import configuration
# ------------------------------------------------------------

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000

SUPPORTED_IMPORT_EXTENSIONS = {".csv", ".xlsx"}

IMPORT_FIELD_LIMITS = {
    "business_name": 250,
    "registration_number": 100,
    "tax_identifier": 100,
    "phone_number": 30,
    "email": 255,
    "address": 500,
    "city": 100,
    "state": 100,
    "source_name": 200,
}


def normalize_csv_header(value: str) -> str:
    value = value.strip().lower()
    return re.sub(r"[^a-z0-9]+", "_", value).strip("_")


def normalize_business_name(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).casefold()
    return re.sub(r"[^\w]+", " ", value, flags=re.UNICODE).strip()


def normalize_identifier(value: str | None) -> str | None:
    if not value:
        return None

    return re.sub(r"[^a-z0-9]", "", value.casefold()) or None


def normalize_phone(value: str | None) -> str | None:
    if not value:
        return None

    digits = re.sub(r"\D", "", value)
    return digits or None


def years_overlap(first: int | None, second: int | None) -> bool:
    # Unknown years are treated conservatively during duplicate checks.
    return first is None or second is None or first == second


def is_duplicate_business(candidate: dict, existing: dict) -> str | None:
    if not years_overlap(candidate["tax_year"], existing["tax_year"]):
        return None

    for field in ("registration_number", "tax_identifier"):
        new_value = normalize_identifier(candidate.get(field))
        old_value = normalize_identifier(existing.get(field))

        if new_value and old_value and new_value == old_value:
            return f"matching {field}"

    same_name = (
        candidate["normalized_name"] == existing["normalized_name"]
    )

    same_city = (
        (candidate.get("city") or "").strip().casefold()
        == (existing.get("city") or "").strip().casefold()
    )

    if same_name and same_city:
        new_phone = normalize_phone(candidate.get("phone_number"))
        old_phone = normalize_phone(existing.get("phone_number"))

        new_email = (candidate.get("email") or "").strip().casefold()
        old_email = (existing.get("email") or "").strip().casefold()

        if new_phone and old_phone and new_phone == old_phone:
            return "matching business name, city, and phone"

        if new_email and old_email and new_email == old_email:
            return "matching business name, city, and email"

    return None


# ------------------------------------------------------------
# Shared CSV/XLSX parsing
# ------------------------------------------------------------

def cell_to_text(value) -> str:
    """
    Convert CSV/Excel cell values into trimmed strings.

    Excel dates, numbers and other non-string values are converted
    to text so the common validation logic can process them.
    """
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return str(value).strip()


def read_import_rows(raw: bytes, extension: str) -> tuple[list[str], list[tuple[int, dict]]]:
    """
    Return normalized headers and rows in a common format:

        (headers, [(spreadsheet_row_number, normalized_row_dict), ...])

    Both CSV and XLSX use the same validation and duplicate-detection
    logic after this function.
    """
    if extension == ".csv":
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=400,
                detail="CSV must use UTF-8 encoding.",
            ) from exc

        try:
            reader = csv.reader(io.StringIO(content, newline=""))

            try:
                original_headers = next(reader)
            except StopIteration:
                raise HTTPException(
                    status_code=400,
                    detail="CSV is empty or has no header row.",
                )

            headers = [
                normalize_csv_header(cell_to_text(header))
                for header in original_headers
            ]

            if not headers or not any(headers):
                raise HTTPException(
                    status_code=400,
                    detail="CSV is empty or has no valid header row.",
                )

            validate_import_headers(headers, "CSV")

            parsed_rows = []

            # CSV reader.line_num accounts for the physical line number,
            # including records that contain quoted line breaks.
            for values in reader:
                row_number = reader.line_num

                # Skip completely empty records.
                if not any(cell_to_text(value) for value in values):
                    continue

                if len(values) > len(headers):
                    # Extra values cannot be mapped safely to a header.
                    parsed_rows.append(
                        (
                            row_number,
                            {
                                "__row_error__": (
                                    "This row contains more values than "
                                    "the header row."
                                )
                            },
                        )
                    )
                    continue

                values = values + [""] * (len(headers) - len(values))

                row = {
                    headers[index]: cell_to_text(value)
                    for index, value in enumerate(values)
                }

                parsed_rows.append((row_number, row))

            return headers, parsed_rows

        except csv.Error as exc:
            raise HTTPException(
                status_code=400,
                detail=f"Could not parse CSV: {exc}",
            ) from exc

    if extension == ".xlsx":
        try:
            workbook = load_workbook(
                io.BytesIO(raw),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=(
                    "Could not read the Excel file. Please upload a valid "
                    ".xlsx workbook."
                ),
            ) from exc

        try:
            worksheet = workbook.active

            if worksheet.max_row is None or worksheet.max_row < 1:
                raise HTTPException(
                    status_code=400,
                    detail="Excel worksheet is empty or has no header row.",
                )

            first_row = next(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=1,
                    values_only=True,
                ),
                None,
            )

            if not first_row:
                raise HTTPException(
                    status_code=400,
                    detail="Excel worksheet has no header row.",
                )

            original_headers = list(first_row)

            # Ignore trailing blank header cells.
            while original_headers and not cell_to_text(original_headers[-1]):
                original_headers.pop()

            headers = [
                normalize_csv_header(cell_to_text(header))
                for header in original_headers
            ]

            if not headers or not any(headers):
                raise HTTPException(
                    status_code=400,
                    detail="Excel worksheet has no valid header row.",
                )

            validate_import_headers(headers, "Excel")

            parsed_rows = []

            for row_number, values in enumerate(
                worksheet.iter_rows(
                    min_row=2,
                    max_col=len(headers),
                    values_only=True,
                ),
                start=2,
            ):
                values = list(values)

                if not any(cell_to_text(value) for value in values):
                    continue

                values = values + [""] * (len(headers) - len(values))

                row = {
                    headers[index]: cell_to_text(value)
                    for index, value in enumerate(values[:len(headers)])
                }

                parsed_rows.append((row_number, row))

            return headers, parsed_rows

        finally:
            workbook.close()

    raise HTTPException(
        status_code=400,
        detail="Unsupported file type. Upload a .csv or .xlsx file.",
    )


def validate_import_headers(headers: list[str], file_type: str) -> None:
    if any(not header for header in headers):
        raise HTTPException(
            status_code=400,
            detail=(
                f"{file_type} contains a blank header. "
                "Every column must have a unique, non-empty header."
            ),
        )

    if len(headers) != len(set(headers)):
        raise HTTPException(
            status_code=400,
            detail=(
                f"{file_type} contains duplicate headers after "
                "normalization. Please rename the duplicate columns."
            ),
        )

    if "business_name" not in headers:
        raise HTTPException(
            status_code=400,
            detail=(
                f"{file_type} must include a business_name column. "
                "For example, use the header Business Name."
            ),
        )


# ------------------------------------------------------------
# CSV and Excel import endpoint
# ------------------------------------------------------------

@router.post("/import-csv")
async def import_businesses_csv(
    file: UploadFile = File(...),
    dry_run: bool = Query(default=False),
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    filename = file.filename or "upload"
    extension = (
        "." + filename.rsplit(".", 1)[-1].lower()
        if "." in filename
        else ""
    )

    if extension not in SUPPORTED_IMPORT_EXTENSIONS:
        await file.close()
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. Please upload a CSV (.csv) "
                "or Excel workbook (.xlsx)."
            ),
        )

    try:
        raw = await file.read(MAX_UPLOAD_BYTES + 1)
    finally:
        await file.close()

    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail="The uploaded file exceeds the 5 MB limit.",
        )

    headers, parsed_rows = read_import_rows(raw, extension)

    if len(parsed_rows) > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=413,
            detail=(
                f"The file exceeds the {MAX_IMPORT_ROWS}-row limit. "
                "Please split the file into smaller batches."
            ),
        )

    prepared = []
    invalid_rows = []

    for row_number, row in parsed_rows:
        if "__row_error__" in row:
            invalid_rows.append({
                "row": row_number,
                "errors": [row["__row_error__"]],
            })
            continue

        if not any(value for value in row.values()):
            continue

        errors = []
        business_name = row.get("business_name", "").strip()

        if not business_name:
            errors.append("business_name is required.")
        elif len(business_name) > IMPORT_FIELD_LIMITS["business_name"]:
            errors.append("business_name exceeds 250 characters.")

        data = {}

        for field, max_length in IMPORT_FIELD_LIMITS.items():
            value = row.get(field, "").strip()
            data[field] = value or None

            if value and len(value) > max_length:
                errors.append(
                    f"{field} exceeds {max_length} characters."
                )

        tax_year_text = row.get("tax_year", "").strip()
        tax_year = None

        if tax_year_text:
            try:
                # Accept Excel numeric values such as 2026.0 as well
                # as text values such as "2026".
                parsed_year = float(tax_year_text)

                if not parsed_year.is_integer():
                    errors.append("tax_year must be a whole number.")
                else:
                    tax_year = int(parsed_year)

                    if not 2000 <= tax_year <= 2100:
                        errors.append(
                            "tax_year must be between 2000 and 2100."
                        )
            except ValueError:
                errors.append("tax_year must be a whole number.")

        source_url = row.get("source_url", "").strip() or None

        if source_url and len(source_url) > 5000:
            errors.append("source_url exceeds 5000 characters.")

        if errors:
            invalid_rows.append({
                "row": row_number,
                "errors": errors,
            })
            continue

        data["business_name"] = business_name
        data["normalized_name"] = normalize_business_name(business_name)
        data["tax_year"] = tax_year
        data["source_url"] = source_url
        data["source_name"] = data.get("source_name") or filename
        data["state"] = data.get("state") or "Edo"

        prepared.append({
            "row": row_number,
            "data": data,
        })

    # --------------------------------------------------------
    # Nothing valid to import
    # --------------------------------------------------------

    if not prepared:
        return {
            "dry_run": dry_run,
            "filename": filename,
            "imported": 0,
            "would_import": 0 if dry_run else None,
            "duplicates_skipped": 0,
            "duplicate_rows": [],
            "duplicate_rows_truncated": False,
            "invalid_rows_count": len(invalid_rows),
            "invalid_rows": invalid_rows[:100],
            "invalid_rows_truncated": len(invalid_rows) > 100,
        }

    # --------------------------------------------------------
    # Find potentially existing records
    # --------------------------------------------------------

    candidate_names = {
        item["data"]["normalized_name"]
        for item in prepared
    }

    registration_numbers = {
        item["data"]["registration_number"]
        for item in prepared
        if item["data"].get("registration_number")
    }

    tax_identifiers = {
        item["data"]["tax_identifier"]
        for item in prepared
        if item["data"].get("tax_identifier")
    }

    conditions = [
        BusinessRegister.normalized_name.in_(candidate_names)
    ]

    if registration_numbers:
        conditions.append(
            BusinessRegister.registration_number.in_(
                registration_numbers
            )
        )

    if tax_identifiers:
        conditions.append(
            BusinessRegister.tax_identifier.in_(tax_identifiers)
        )

    existing_rows = (
        db.query(BusinessRegister)
        .filter(or_(*conditions))
        .all()
    )

    def comparison_record(business):
        return {
            "registration_number": business.registration_number,
            "tax_identifier": business.tax_identifier,
            "normalized_name": business.normalized_name,
            "city": business.city,
            "phone_number": business.phone_number,
            "email": business.email,
            "tax_year": business.tax_year,
        }

    known_records = [
        comparison_record(item)
        for item in existing_rows
    ]

    accepted = []
    duplicates = []

    # --------------------------------------------------------
    # Detect duplicates against the database and this upload
    # --------------------------------------------------------

    for item in prepared:
        data = item["data"]
        duplicate_reason = None

        for known in known_records:
            duplicate_reason = is_duplicate_business(data, known)

            if duplicate_reason:
                break

        if duplicate_reason:
            duplicates.append({
                "row": item["row"],
                "reason": duplicate_reason,
            })
            continue

        accepted.append(BusinessRegister(**data))

        # Prevent duplicate records within the same uploaded file.
        known_records.append(data)

    # --------------------------------------------------------
    # Dry run: preview without writing to the database
    # --------------------------------------------------------

    if dry_run:
        return {
            "dry_run": True,
            "filename": filename,
            "would_import": len(accepted),
            "duplicates_skipped": len(duplicates),
            "duplicate_rows": duplicates[:100],
            "duplicate_rows_truncated": len(duplicates) > 100,
            "invalid_rows_count": len(invalid_rows),
            "invalid_rows": invalid_rows[:100],
            "invalid_rows_truncated": len(invalid_rows) > 100,
        }

    # --------------------------------------------------------
    # Commit import
    # --------------------------------------------------------

    try:
        db.add_all(accepted)
        db.commit()
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail=(
                "Import failed. The database transaction was rolled back."
            ),
        ) from exc

    return {
        "dry_run": False,
        "filename": filename,
        "imported": len(accepted),
        "duplicates_skipped": len(duplicates),
        "duplicate_rows": duplicates[:100],
        "duplicate_rows_truncated": len(duplicates) > 100,
        "invalid_rows_count": len(invalid_rows),
        "invalid_rows": invalid_rows[:100],
        "invalid_rows_truncated": len(invalid_rows) > 100,
    }


# ------------------------------------------------------------
# Manual match review
# ------------------------------------------------------------

@router.post("/{business_id}/match-review")
def review_business_match(
    business_id: int,
    data: BusinessMatchReviewRequest,
    db: Session = Depends(get_db),
    _: User = Depends(require_admin),
):
    action = data.action.strip().upper()

    if action not in {"CONFIRM", "REJECT"}:
        raise HTTPException(
            status_code=400,
            detail="Action must be CONFIRM or REJECT.",
        )

    business = (
        db.query(BusinessRegister)
        .filter(BusinessRegister.id == business_id)
        .first()
    )

    if business is None:
        raise HTTPException(
            status_code=404,
            detail="Business register record not found.",
        )

    if action == "CONFIRM":
        if data.taxpayer_id is None:
            raise HTTPException(
                status_code=400,
                detail="taxpayer_id is required to confirm a match.",
            )

        taxpayer = (
            db.query(Taxpayer)
            .filter(Taxpayer.id == data.taxpayer_id)
            .first()
        )

        if taxpayer is None:
            raise HTTPException(
                status_code=404,
                detail="Taxpayer not found.",
            )

        business.matched_taxpayer_id = taxpayer.id
        business.match_status = "CONFIRMED"

    else:
        business.matched_taxpayer_id = None
        business.match_confidence = None
        business.match_status = "REJECTED"

    if data.review_notes is not None:
        business.review_notes = data.review_notes.strip() or None

    try:
        db.commit()
        db.refresh(business)
    except Exception as exc:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="Match review failed; changes were rolled back.",
        ) from exc

    result = serialize_business(business)
    result["matched_taxpayer"] = None

    if business.matched_taxpayer_id is not None:
        taxpayer = (
            db.query(Taxpayer)
            .filter(Taxpayer.id == business.matched_taxpayer_id)
            .first()
        )

        if taxpayer is not None:
            result["matched_taxpayer"] = {
                "id": taxpayer.id,
                "reference": getattr(
                    taxpayer,
                    "taxpayer_reference",
                    None,
                ),
                "full_name": taxpayer.full_name,
            }

    return result